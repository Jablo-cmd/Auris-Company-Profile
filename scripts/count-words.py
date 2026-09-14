#!/usr/bin/env python3
"""
Authoritative Stage 10 word count.

Counts visible reader-facing text in the generated .docx: paragraphs and table
cells, excluding headers/footers (python-docx never exposes these as body
paragraphs), the Contents-page field/instruction labels, and the
"confirmation pending" marker phrase (it stands in for absent content, not
authored copy, so it is reported separately rather than folded into
words_actual).

See 01_Project/09_Stage10-Production-Methodology.md for the full rule set.

Usage: python3 scripts/count-words.py [path/to/file.docx]
"""
import sys
import os
import docx

DEFAULT_PATH = os.path.join(os.path.dirname(__file__), "..", "08_Word",
                             "Auris-Nexus-Company-Profile-Stage10-Draft.docx")
PENDING_MARK = "confirmation pending"
EXCLUDE_TEXT_EXACT = {
    "Contents",
    'Auto-generated field. In Word: right-click → "Update Field" (or Update Table) after opening.',
}


def count(path):
    d = docx.Document(path)
    total_words = 0
    pending_phrase_words = 0

    for p in d.paragraphs:
        t = p.text.strip()
        if not t or t in EXCLUDE_TEXT_EXACT:
            continue
        total_words += len(t.split())
        pending_phrase_words += t.count(PENDING_MARK) * len(PENDING_MARK.split())

    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if not txt:
                    continue
                total_words += len(txt.split())
                pending_phrase_words += txt.count(PENDING_MARK) * len(PENDING_MARK.split())

    authored = total_words - pending_phrase_words
    return total_words, pending_phrase_words, authored


if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PATH
    total, pending, authored = count(path)
    print(f"File: {path}")
    print(f"Raw visible word count (cover + body + tables):  {total}")
    print(f"Words contributed by 'confirmation pending':     {pending}")
    print(f"AUTHORED word count (authoritative words_actual): {authored}")
